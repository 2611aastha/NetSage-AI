import os
import zipfile
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 1. Initialize Document & Set Page Margins for Perfect 1-Page Layout
doc = docx.Document()
for section in doc.sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# Helper functions for styling
def add_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(17)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(15, 23, 42)

def add_meta():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    items = [
        ("Author: ", "Aastha Deep"),
        ("  |  Roll No: ", "102303283"),
        ("  |  Email: ", "adeep_be23@thapar.edu"),
        ("  |  Institution: ", "Thapar Institute")
    ]
    for label, val in items:
        r1 = p.add_run(label)
        r1.font.size = Pt(9)
        r1.font.name = 'Arial'
        r1.font.color.rgb = RGBColor(100, 116, 139)
        r2 = p.add_run(val)
        r2.bold = True
        r2.font.size = Pt(9)
        r2.font.name = 'Arial'
        r2.font.color.rgb = RGBColor(30, 41, 59)

def add_heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(29, 78, 216)

def add_body(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(9)
        r_pre.font.name = 'Arial'
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.name = 'Arial'
    r.font.color.rgb = RGBColor(51, 65, 85)

# --- Build Document Content ---
add_title("NetSage AI — Evidence-Backed Network Troubleshooting Dashboard")
add_meta()

add_heading("1. Executive Summary & Project Objective")
add_body("NetSage AI is an evidence-backed network diagnostic and decision-support platform built for Cisco Packet Tracer environments. It bridges deterministic rule validation with AI-driven troubleshooting while enforcing strict Responsible AI principles—guaranteeing that AI never executes CLI commands automatically and requiring explicit human operator review before taking action.")

add_heading("2. System Architecture & Tech Stack")
add_body("FastAPI (Python 3.13) serving asynchronous REST API endpoints for rule checking, AI diagnostics, and audit logging.", "Backend Engine: ")
add_body("React 18 with Vite, Tailwind CSS, and Lucide React icons delivering a high-density, dark-mode network engineering dashboard.", "Frontend Dashboard: ")
add_body("Custom regex engine analyzing Cisco CLI evidence ('show ip route', 'show interface') deterministically before AI invocation.", "Rule Engine: ")
add_body("FastAPI backend orchestrates rule checks and queries LLM providers with full offline mock fallback for 100% demo reliability.", "Diagnostic Pipeline: ")

add_heading("3. Key Requirements Met & Implementation Details")
add_body("Integrated 32 pre-configured synthetic cases covering IP mismatches, administratively down interfaces, subnet errors, and missing default gateways.", "32 Synthetic Test Cases: ")
add_body("Rule-checker guarantees instant pass/fail validation on physical/link layers, preventing AI hallucinations.", "Deterministic First Approach: ")
add_body("UI forces network engineers to explicitly Accept or Reject AI suggestions, logging decisions to data/audit_log.csv.", "Mandatory Human Review Gate: ")
add_body("Includes documented human corrections in responsible_ai_log.csv to demonstrate advisory boundaries and human oversight.", "Responsible AI Audit Trail: ")

add_heading("4. Testing & Submission Deliverables")
add_body("Automated test suite implemented in tests/test_checker.py and verified with pytest to guarantee zero regressions.", "Unit Testing (Pytest): ")
add_body("Standalone Python packaging script automates clean source code archiving into NetSage_AI_Submission.zip.", "Packaging Automation: ")

doc.save("NetSage_AI_Project_Summary.docx")
print("NetSage_AI_Project_Summary.docx created successfully!")

# 2. Package Clean Source Tree into ZIP
zip_name = "NetSage_AI_Submission.zip"
with zipfile.ZipFile(zip_name, 'w', zipfile.ZIP_DEFLATED) as zipf:
    for root, dirs, files in os.walk("."):
        if any(ignored in root for ignored in ["venv", ".git", "node_modules", "__pycache__"]):
            continue
        for file in files:
            if file not in [zip_name, "NetSage_AI_Project_Summary.docx", "package_project.py"]:
                filepath = os.path.join(root, file)
                zipf.write(filepath, os.path.relpath(filepath, "."))

print("NetSage_AI_Submission.zip created successfully.")